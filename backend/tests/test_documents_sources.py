import hashlib
import io
import zipfile

import pymupdf
from docx import Document

from backend.app.database import ROOT
from backend.app.documents import extract_document


def test_source_evidence_and_ledger(client):
    value = {'title': 'Fuente de prueba', 'url': 'https://example.org/publico', 'description': 'Registro de prueba automatizada', 'observed_at': '2026-10-24', 'notes': '=No ejecutar fórmula'}
    response = client.post('/api/sources', json=value)
    assert response.status_code == 200
    source_id = response.json()['id']
    content = b'Evidencia local de prueba'
    response = client.post(f'/api/sources/{source_id}/evidence', files={'file': ('evidencia.txt', content, 'text/plain')})
    assert response.status_code == 200
    assert response.json()['sha256'] == hashlib.sha256(content).hexdigest()
    assert client.get('/api/files/' + response.json()['id']).content == content
    source = client.get('/api/sources').json()['items'][0]
    assert len(source['evidence']) == 1
    exported = client.get('/api/sources/export').text
    assert hashlib.sha256(content).hexdigest() in exported
    assert "'=No ejecutar" in exported
    assert client.post('/api/sources', json={**value, 'url': 'javascript:alert(1)'}).status_code == 422
    assert client.post('/api/sources', json={**value, 'url': 'http://127.0.0.1/private'}).status_code == 422


def test_review_required_export_and_invalidation(client):
    dossier = client.post('/api/dossiers/demo').json()
    id_ = dossier['id']
    assert dossier['demo'] and not dossier['reviewed']
    assert client.get(f'/api/dossiers/{id_}/export/pdf').status_code == 409
    update = client.put(f'/api/dossiers/{id_}', json={'fields': dossier['fields'], 'reviewed': True})
    assert update.status_code == 200
    pdf = client.get(f'/api/dossiers/{id_}/export/pdf')
    assert pdf.status_code == 200
    with pymupdf.open(stream=pdf.content, filetype='pdf') as doc:
        text = ''.join(p.get_text() for p in doc)
        assert 'FICTICIA' in text and 'Persona de demostración' in text
    word = client.get(f'/api/dossiers/{id_}/export/docx')
    doc = Document(io.BytesIO(word.content))
    assert any('FICTICIA' in p.text for p in doc.paragraphs)
    edited = {**dossier['fields'], 'name': 'Otro nombre ficticio'}
    assert not client.put(f'/api/dossiers/{id_}', json={'fields': edited}).json()['reviewed']
    assert client.get(f'/api/dossiers/{id_}/export/docx').status_code == 409


def test_all_six_original_identification_documents():
    with zipfile.ZipFile(ROOT / 'hakathon-20261024T164620Z-1-001.zip') as z:
        names = [n for n in z.namelist() if '/automatizacion de fichas' in n and n.endswith(('.docx', '.pdf'))]
        assert len(names) == 6
        for n in names:
            text, fields = extract_document(n, z.read(n))
            assert len(text) > 100
            assert 'name' in fields and 'notes' in fields


def test_import_edit_missing_and_empty_pdf(client):
    document = Document()
    document.add_paragraph('Nombre: Persona ficticia de prueba')
    document.add_paragraph('Fecha de nacimiento: 1990-01-01')
    buffer = io.BytesIO()
    document.save(buffer)
    response = client.post('/api/dossiers/upload', files={'file': ('prueba.docx', buffer.getvalue())})
    assert response.status_code == 200
    dossier = response.json()
    assert dossier['fields']['name'] == 'Persona ficticia de prueba'
    assert 'Domicilio' in dossier['missing']
    assert not dossier['reviewed']
    pdf = pymupdf.open()
    pdf.new_page()
    empty = client.post('/api/dossiers/upload', files={'file': ('escaneo.pdf', pdf.tobytes())})
    assert empty.status_code == 422
    assert 'OCR' in empty.json()['detail']
