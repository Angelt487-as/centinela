import io
import re
from pathlib import Path
from xml.sax.saxutils import escape

import pymupdf as fitz
from docx import Document
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor
from fastapi import HTTPException
from reportlab.lib import colors
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.lib.enums import TA_LEFT
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer

from .parsers import key

FIELD_LABELS = {
    'name': 'Nombre completo', 'alias': 'Alias', 'birth_date': 'Fecha de nacimiento',
    'birth_place': 'Lugar de nacimiento', 'sex': 'Sexo', 'nationality': 'Nacionalidad',
    'curp': 'CURP', 'address': 'Domicilio', 'occupation': 'Ocupación', 'notes': 'Observaciones',
}
MATCH_LABELS = {
    'name': ['nombre completo', 'nombre'], 'alias': ['alias', 'apodo'],
    'birth_date': ['fecha de nacimiento', 'fecha nacimiento'],
    'birth_place': ['lugar de nacimiento', 'lugar nacimiento'], 'sex': ['sexo'],
    'nationality': ['nacionalidad'], 'curp': ['curp'],
    'address': ['domicilio', 'dirección', 'direccion'], 'occupation': ['ocupación', 'ocupacion'],
    'notes': ['observaciones'],
}


def extract_document(name, content):
    try:
        if Path(name).suffix.lower() == '.pdf':
            with fitz.open(stream=content, filetype='pdf') as doc:
                if doc.needs_pass:
                    raise HTTPException(422, 'El PDF tiene contraseña. Importa una copia autorizada sin contraseña.')
                if len(doc) > 250:
                    raise HTTPException(413, 'El documento supera el límite de 250 páginas.')
                text = '\n'.join(page.get_text(sort=True) for page in doc)
        else:
            doc = Document(io.BytesIO(content))
            text = '\n'.join(''.join(t.text or '' for t in p.iter(qn('w:t')))
                             for p in doc.element.iter(qn('w:p')))
        if len(text) > 2_000_000:
            raise HTTPException(413, 'El documento contiene demasiado texto para este MVP.')
        if not text.strip():
            raise HTTPException(422, 'No hay texto extraíble. Este MVP no incluye OCR para documentos escaneados.')
    except HTTPException:
        raise
    except Exception:
        raise HTTPException(422, 'No se pudo extraer el texto. Verifica que el documento no esté dañado.') from None
    fields = {k: '' for k in FIELD_LABELS}
    lines = [line.strip() for line in text.splitlines() if line.strip()]
    all_labels = {key(label) for labels in MATCH_LABELS.values() for label in labels}
    for field, labels in MATCH_LABELS.items():
        for i, line in enumerate(lines):
            matched = False
            for label in labels:
                match = re.match(r'^' + re.escape(label) + r'\s*[:：\-]\s*(.+)$', line, re.I)
                if match and len(match.group(1)) <= 1000:
                    fields[field] = match.group(1).strip()
                    matched = True
                    break
                if key(line.rstrip(':')) == key(label) and i + 1 < len(lines):
                    candidate = lines[i + 1]
                    if key(candidate.rstrip(':')) not in all_labels and len(candidate) <= 1000:
                        fields[field] = candidate
                        matched = True
                        break
            if matched:
                break
    return text, fields


def export_dossier(row, fields, format_):
    banner = 'DEMOSTRACIÓN FICTICIA · SIN VALIDEZ DOCUMENTAL' if row['demo'] else 'Revisión humana confirmada'
    if format_ == 'docx':
        doc = Document()
        section = doc.sections[0]
        section.top_margin = Inches(.8)
        section.bottom_margin = Inches(.8)
        style = doc.styles['Normal']
        style.font.name = 'Arial'
        style.font.size = Pt(11)
        for name in ('Title', 'Heading 1', 'Heading 2'):
            doc.styles[name].font.name = 'Arial'
            doc.styles[name].font.color.rgb = RGBColor(0, 0, 0)
        doc.add_heading('Centinela 360', 0)
        doc.add_heading('Ficha de identificación', 1)
        doc.add_paragraph(banner)
        for field, label in FIELD_LABELS.items():
            doc.add_heading(label, 2)
            doc.add_paragraph(fields.get(field) or 'No proporcionado')
        doc.add_paragraph('Documento generado a partir de información revisada. Los campos ausentes no se infieren.')
        buffer = io.BytesIO()
        doc.save(buffer)
    else:
        buffer = io.BytesIO()
        pdf = SimpleDocTemplate(buffer, title='Ficha de identificación Centinela 360',
                                author='Centinela 360', leftMargin=48, rightMargin=48,
                                topMargin=48, bottomMargin=48)
        styles = getSampleStyleSheet()
        styles.add(ParagraphStyle(name='FieldValue', fontName='Helvetica', fontSize=11, leading=16,
                                  spaceAfter=13, wordWrap='CJK'))
        story = [Paragraph('Centinela 360', styles['Title']),
                 Paragraph('Ficha de identificación', styles['Heading1']),
                 Paragraph(escape(banner), styles['Normal']), Spacer(1, 18)]
        for field, label in FIELD_LABELS.items():
            value = re.sub(r'[\x00-\x08\x0b-\x1f]', '', fields.get(field) or 'No proporcionado')
            story.extend([Paragraph(escape(label), styles['Heading2']),
                          Paragraph(escape(value).replace('\n', '<br/>'), styles['FieldValue'])])
        story.extend([Spacer(1, 12), Paragraph('Los campos ausentes no se infieren.', styles['Normal'])])
        pdf.build(story)
    return buffer.getvalue()
